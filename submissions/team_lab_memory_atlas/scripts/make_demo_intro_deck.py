from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "presentation_assets"
OUT = ROOT / "demo" / "lab_memory_atlas_intro.pptx"

NAVY = RGBColor(23, 35, 45)
NAVY_2 = RGBColor(31, 50, 63)
SURFACE = RGBColor(243, 246, 248)
PANEL = RGBColor(255, 255, 255)
INK = RGBColor(24, 34, 43)
MUTED = RGBColor(92, 102, 112)
ORANGE = RGBColor(245, 130, 32)
TEAL = RGBColor(11, 111, 106)
LINE = RGBColor(215, 221, 226)


def rgb(hex_value: str) -> RGBColor:
    value = hex_value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def fill(shape, color: RGBColor, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def line(shape, color: RGBColor, width: float = 1, transparency: int = 0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def no_line(shape) -> None:
    shape.line.fill.background()


def text_box(slide, text: str, x: float, y: float, w: float, h: float, size: int,
             color: RGBColor = INK, bold: bool = False, font: str = "Aptos",
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    frame.text = text
    for p in frame.paragraphs:
        p.alignment = align
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def multiline(slide, lines: list[tuple[str, str]], x: float, y: float, w: float, h: float,
              body_color: RGBColor = MUTED):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = 0
    frame.margin_bottom = 0
    for idx, (head, body) in enumerate(lines):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.space_after = Pt(8)
        r1 = p.add_run()
        r1.text = head
        r1.font.name = "Aptos"
        r1.font.size = Pt(16)
        r1.font.bold = True
        r1.font.color.rgb = INK
        r2 = p.add_run()
        r2.text = f" {body}"
        r2.font.name = "Aptos"
        r2.font.size = Pt(16)
        r2.font.color.rgb = body_color
    return box


def add_brand(slide, dark: bool = False) -> None:
    logo = ASSET_DIR / "knight-logo.png"
    if logo.exists():
        slide.shapes.add_picture(str(logo), Inches(0.58), Inches(0.42), height=Inches(0.52))
    color = PANEL if dark else INK
    sub = RGBColor(196, 209, 209) if dark else MUTED
    text_box(slide, "de ridder lab", 1.18, 0.43, 1.9, 0.22, 11, color=color, bold=True)
    text_box(slide, "Lab Memory Atlas", 1.18, 0.68, 1.9, 0.18, 8, color=sub)


def add_pill(slide, text: str, x: float, y: float, w: float, fill_color: RGBColor,
             text_color: RGBColor = PANEL) -> None:
    pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    fill(pill, fill_color)
    no_line(pill)
    tf = pill.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = text_color


def slide_one(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    add_brand(slide, dark=True)

    text_box(slide, "Why we made this", 0.62, 1.12, 2.7, 0.22, 11, color=ORANGE, bold=True)
    text_box(slide, "Lab knowledge disappears\nafter meetings", 0.58, 1.48, 6.0, 1.58, 36, color=PANEL, bold=True, font="Georgia")
    text_box(
        slide,
        "Meetings create decisions, ideas, and follow-up tasks. Without a shared memory layer, they become hard to find, hard to follow, and hard to connect.",
        0.62,
        3.52,
        5.95,
        0.8,
        16,
        color=RGBColor(214, 224, 224),
    )

    card_data = [
        ("Ideas", "Good suggestions stay in the room."),
        ("Actions", "Tasks are not assigned, revisited, or closed."),
        ("Connections", "People miss who can help with what."),
    ]
    for i, (title, body) in enumerate(card_data):
        x = 0.62 + i * 2.15
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.12), Inches(1.84), Inches(1.08))
        fill(card, NAVY_2)
        line(card, rgb("#405667"), 1)
        text_box(slide, title, x + 0.16, 5.28, 1.4, 0.22, 13, color=ORANGE, bold=True)
        text_box(slide, body, x + 0.16, 5.62, 1.48, 0.42, 10, color=RGBColor(214, 224, 224))

    # Right-side meeting memory diagram.
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.25), Inches(1.08), Inches(5.42), Inches(5.38))
    fill(panel, rgb("#20313c"))
    line(panel, rgb("#4b6272"), 1)
    text_box(slide, "The gap", 7.55, 1.38, 1.2, 0.24, 12, color=ORANGE, bold=True)
    text_box(slide, "Useful context exists, but it is not structured.", 7.55, 1.7, 3.7, 0.36, 12, color=RGBColor(214, 224, 224))

    doc = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.62), Inches(2.48), Inches(1.45), Inches(1.5))
    fill(doc, PANEL)
    line(doc, rgb("#b4c0c8"), 1)
    text_box(slide, "Transcript", 7.84, 2.72, 0.92, 0.22, 10, color=INK, bold=True, align=PP_ALIGN.CENTER)
    for offset in [3.04, 3.25, 3.46]:
        segment = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(7.86), Inches(offset), Inches(0.98), Inches(0.05))
        fill(segment, LINE)
        no_line(segment)

    cloud = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CLOUD, Inches(9.52), Inches(2.42), Inches(1.6), Inches(1.22))
    fill(cloud, rgb("#f7d0b2"), 20)
    line(cloud, ORANGE, 1)
    text_box(slide, "memory gap", 9.82, 2.9, 0.95, 0.22, 10, color=RGBColor(214, 224, 224), bold=True, align=PP_ALIGN.CENTER)

    target = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(11.52), Inches(2.48), Inches(0.58), Inches(1.5))
    fill(target, TEAL)
    no_line(target)
    text_box(slide, "?", 11.68, 2.78, 0.25, 0.45, 28, color=PANEL, bold=True, align=PP_ALIGN.CENTER)

    for x1, x2 in [(9.12, 9.5), (11.15, 11.5)]:
        arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(3.2), Inches(x2), Inches(3.2))
        arrow.line.color.rgb = RGBColor(214, 224, 224)
        arrow.line.width = Pt(2)
        arrow.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    add_pill(slide, "decisions", 7.7, 4.82, 1.08, ORANGE)
    add_pill(slide, "tasks", 9.05, 4.82, 0.78, TEAL)
    add_pill(slide, "helpers", 10.08, 4.82, 0.92, rgb("#a9442e"))
    add_pill(slide, "progress", 11.23, 4.82, 0.98, rgb("#d6a21d"), INK)


def slide_growth(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = SURFACE
    add_brand(slide, dark=False)

    text_box(slide, "Why this gets worse over time", 0.62, 1.06, 3.2, 0.22, 11, color=ORANGE, bold=True)
    text_box(slide, "As the lab grows, informal memory stops scaling", 0.58, 1.38, 6.2, 1.08, 32, color=INK, bold=True, font="Georgia")
    text_box(
        slide,
        "More people, SIGs, and projects are a strength. But without a shared map, it becomes harder to keep cohesion: who is doing what, which ideas are active, and where progress is stuck.",
        0.62,
        2.58,
        5.95,
        0.76,
        14,
        color=MUTED,
    )

    for i, (number, label, color) in enumerate(
        [
            ("29", "people to keep in view", TEAL),
            ("11", "active groups and projects", ORANGE),
            ("3", "SIGs with overlapping work", rgb("#a9442e")),
        ]
    ):
        x = 0.62 + i * 1.9
        metric = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.68), Inches(1.55), Inches(1.02))
        fill(metric, PANEL)
        line(metric, rgb("#d7dde2"), 1)
        text_box(slide, number, x + 0.16, 3.82, 0.64, 0.34, 24, color=color, bold=True)
        text_box(slide, label, x + 0.16, 4.23, 1.12, 0.26, 8, color=MUTED)

    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.86), Inches(0.96), Inches(5.7), Inches(5.92))
    fill(panel, PANEL)
    line(panel, rgb("#d0d8de"), 1)
    text_box(slide, "The coordination gap", 7.16, 1.26, 2.4, 0.28, 14, color=INK, bold=True)
    text_box(slide, "Activity rises faster than shared context.", 7.16, 1.58, 2.72, 0.22, 10, color=MUTED)

    # Axes.
    y_axis = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.34), Inches(5.54), Inches(7.34), Inches(2.22))
    y_axis.line.color.rgb = rgb("#aab7c0")
    y_axis.line.width = Pt(1.1)
    x_axis = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.34), Inches(5.54), Inches(11.82), Inches(5.54))
    x_axis.line.color.rgb = rgb("#aab7c0")
    x_axis.line.width = Pt(1.1)
    text_box(slide, "lab size", 11.2, 5.72, 0.72, 0.16, 8, color=MUTED)
    text_box(slide, "context", 7.0, 2.08, 0.58, 0.16, 8, color=MUTED)

    # Rising work/project line.
    rising = [(7.54, 5.18), (8.34, 4.72), (9.16, 4.14), (10.0, 3.24), (11.48, 2.42)]
    for (x1, y1), (x2, y2) in zip(rising, rising[1:]):
        seg = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        seg.line.color.rgb = ORANGE
        seg.line.width = Pt(3)
    text_box(slide, "projects + activity", 10.36, 2.18, 1.2, 0.18, 9, color=ORANGE, bold=True)

    # Falling shared-context line.
    falling = [(7.54, 2.82), (8.34, 3.08), (9.16, 3.54), (10.0, 4.28), (11.48, 5.02)]
    for (x1, y1), (x2, y2) in zip(falling, falling[1:]):
        seg = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        seg.line.color.rgb = TEAL
        seg.line.width = Pt(3)
        seg.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    text_box(slide, "shared context", 10.48, 5.12, 1.1, 0.18, 9, color=TEAL, bold=True)

    gap = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.92), Inches(3.04), Inches(1.58), Inches(0.72))
    fill(gap, NAVY)
    no_line(gap)
    text_box(slide, "cohesion gap", 9.16, 3.22, 1.08, 0.2, 11, color=PANEL, bold=True, align=PP_ALIGN.CENTER)

    for i, (question, detail) in enumerate(
        [
            ("Who owns this?", "tasks and project responsibility"),
            ("What changed?", "progress since the last meeting"),
            ("Who can help?", "expertise across people and SIGs"),
        ]
    ):
        x = 0.62 + i * 2.15
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.58), Inches(1.86), Inches(0.92))
        fill(card, PANEL)
        line(card, rgb("#d7dde2"), 1)
        text_box(slide, question, x + 0.14, 5.75, 1.34, 0.2, 11, color=INK, bold=True)
        text_box(slide, detail, x + 0.14, 6.04, 1.4, 0.22, 8, color=MUTED)


def slide_two(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = SURFACE
    add_brand(slide, dark=False)

    text_box(slide, "What this solves", 0.62, 1.06, 2.5, 0.22, 11, color=ORANGE, bold=True)
    text_box(slide, "A living people and project map for the lab", 0.58, 1.38, 5.85, 0.95, 33, color=INK, bold=True, font="Georgia")
    text_box(
        slide,
        "Meeting transcripts become reviewed lab memory: people, projects, SIGs, supervisors, actions, progress, and suggested helpers.",
        0.62,
        2.44,
        5.82,
        0.66,
        14,
        color=MUTED,
    )

    flow_panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.58), Inches(3.4), Inches(5.55), Inches(2.55))
    fill(flow_panel, PANEL)
    line(flow_panel, LINE, 1)
    multiline(
        slide,
        [
            ("Capture", "meeting transcript and context"),
            ("Draft", "AI proposes summaries, actions, and profile updates"),
            ("Review", "lab members approve or reject changes"),
            ("Connect", "people, projects, supervisors, and follow-ups"),
        ],
        0.92,
        3.74,
        4.9,
        1.86,
    )

    node_specs = [
        ("Transcript", 0.92, 6.28, ORANGE),
        ("AI draft", 2.28, 6.28, TEAL),
        ("Human review", 3.48, 6.28, rgb("#a9442e")),
        ("Lab Atlas", 5.0, 6.28, NAVY),
    ]
    for label, x, y, color in node_specs:
        add_pill(slide, label, x, y, 0.98 if label != "Human review" else 1.22, color)
    for x in [1.9, 3.08, 4.68]:
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(6.45), Inches(x + 0.24), Inches(6.45))
        connector.line.color.rgb = MUTED
        connector.line.width = Pt(1.4)

    # App preview on the right.
    shadow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.94), Inches(1.05), Inches(5.62), Inches(5.72))
    fill(shadow, rgb("#9ba9b3"), 74)
    no_line(shadow)
    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(0.9), Inches(5.62), Inches(5.72))
    fill(frame, PANEL)
    line(frame, rgb("#d0d8de"), 1)
    topbar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6.8), Inches(0.9), Inches(5.62), Inches(0.34))
    fill(topbar, rgb("#edf1f3"))
    no_line(topbar)
    for i, color in enumerate([rgb("#f05f4f"), rgb("#f3b34c"), rgb("#63c174")]):
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(7.04 + i * 0.18), Inches(1.02), Inches(0.08), Inches(0.08))
        fill(dot, color)
        no_line(dot)

    screenshot = ASSET_DIR / "app-home.png"
    if screenshot.exists():
        slide.shapes.add_picture(str(screenshot), Inches(6.92), Inches(1.34), width=Inches(5.38), height=Inches(3.36))

    for i, (head, body) in enumerate(
        [
            ("People", "current work and helpers"),
            ("Projects", "progress and context"),
            ("Actions", "visible to the right people"),
        ]
    ):
        x = 7.04 + i * 1.7
        mini = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.98), Inches(1.48), Inches(0.78))
        fill(mini, rgb("#f7faf9"))
        line(mini, rgb("#dde5e8"), 1)
        text_box(slide, head, x + 0.12, 5.12, 1.16, 0.18, 10, color=TEAL, bold=True)
        text_box(slide, body, x + 0.12, 5.34, 1.18, 0.22, 7, color=MUTED)

    label = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.34), Inches(6.28), Inches(4.3), Inches(0.44))
    fill(label, NAVY)
    no_line(label)
    text_box(slide, "The demo starts from the map, then follows one meeting into reviewed memory.", 7.58, 6.39, 3.8, 0.2, 9, color=PANEL, align=PP_ALIGN.CENTER)


def build() -> Path:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_one(prs)
    slide_growth(prs)
    slide_two(prs)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
